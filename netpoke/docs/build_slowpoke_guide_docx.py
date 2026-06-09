#!/usr/bin/env python3
"""Build slowpoke-paper-deep-dive.docx from slowpoke-paper-deep-dive.md."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
MD = HERE / "slowpoke-paper-deep-dive.md"
OUT = HERE / "slowpoke-paper-deep-dive.docx"


def add_formatted_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            p.add_run(part)


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|?[\s\-:|]+\|?$", line.strip()))


def md_to_docx(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            headers = parse_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Table Grid"
            for j, h in enumerate(headers):
                table.rows[0].cells[j].text = h
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    if c < len(table.rows[r + 1].cells):
                        table.rows[r + 1].cells[c].text = cell
            doc.add_paragraph()
            continue

        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                add_formatted_paragraph(doc, lines[i].strip()[2:], style="List Bullet")
                i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                text = re.sub(r"^\d+\.\s", "", lines[i].strip())
                add_formatted_paragraph(doc, text, style="List Number")
                i += 1
            continue

        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("```")
                or nxt.startswith("|")
                or nxt.startswith("- ")
                or re.match(r"^\d+\.\s", nxt)
                or nxt == "---"
            ):
                break
            para_lines.append(nxt)
            i += 1
        add_formatted_paragraph(doc, " ".join(para_lines))

    doc.save(out_path)
    print(f"Wrote {out_path} ({out_path.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    md = Path(sys.argv[1]) if len(sys.argv) > 1 else MD
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT
    if not md.is_file():
        sys.exit(f"Missing {md}")
    md_to_docx(md, out)
